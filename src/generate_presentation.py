"""
generate_presentation.py
Generates a 12-slide professional academic presentation in PPTX format
(presentation/NLP_Project_Presentation.pptx) using python-pptx.
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import pandas as pd

def set_slide_background(slide, color_rgb):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color_rgb

def add_header(slide, title_text, category_text="NLP CAPSTONE PROJECT"):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = RGBColor(31, 78, 121)
    
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = RGBColor(24, 43, 73)

def add_card(slide, left, top, width, height, bg_rgb=RGBColor(245, 248, 252), border_rgb=RGBColor(210, 222, 235)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_rgb
    shape.line.color.rgb = border_rgb
    shape.line.width = Pt(1.2)
    return shape

def add_table_to_slide(slide, df, left, top, width, height, col_widths=None):
    rows = len(df) + 1
    cols = len(df.columns)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    
    # Headers
    for c_idx, col_name in enumerate(df.columns):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(31, 78, 121)
        p = cell.text_frame.paragraphs[0]
        p.text = str(col_name)
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER
        
    df_reset = df.reset_index(drop=True)
    for r_idx, row in df_reset.iterrows():
        bg = RGBColor(242, 246, 250) if r_idx % 2 == 1 else RGBColor(255, 255, 255)
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            p = cell.text_frame.paragraphs[0]
            p.text = str(val)
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(40, 40, 40)
            if c_idx > 0 and any(char.isdigit() for char in str(val)):
                p.alignment = PP_ALIGN.RIGHT
            else:
                p.alignment = PP_ALIGN.LEFT
                
    if col_widths and len(col_widths) == cols:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
            
    return table_shape

def build_presentation():
    print("[Presentation Generator] Initializing 16:9 Presentation...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank
    
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    fig_dir = os.path.join(base_dir, "figures")
    res_dir = os.path.join(base_dir, "results")
    out_path = os.path.join(base_dir, "presentation", "NLP_Project_Presentation.pptx")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    
    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 1: Title...")
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, RGBColor(24, 43, 73)) # Deep Navy
    
    # Accent banner card
    add_card(s1, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5), 
             bg_rgb=RGBColor(31, 56, 94), border_rgb=RGBColor(60, 95, 145))
    
    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.333), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "SPEECH-TO-TEXT MEETING TRANSCRIPTION AND\nACTION ITEM SUMMARIZER"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = "An End-to-End Deep Learning & Transformer Pipeline (Whisper -> RoBERTa -> spaCy -> BART)"
    p2.font.size = Pt(16)
    p2.font.italic = True
    p2.font.color.rgb = RGBColor(180, 205, 235)
    p2.space_before = Pt(10)
    
    tb_meta = s1.shapes.add_textbox(Inches(1.5), Inches(4.0), Inches(10.333), Inches(2.0))
    tf_m = tb_meta.text_frame
    tf_m.word_wrap = True
    
    pm1 = tf_m.paragraphs[0]
    pm1.text = "Student: [STUDENT NAME]  |  Roll No: [ROLL NO]  |  B.Tech Computer Science & Engineering"
    pm1.font.size = Pt(14)
    pm1.font.color.rgb = RGBColor(240, 240, 240)
    
    pm2 = tf_m.add_paragraph()
    pm2.text = "Supervisor: [GUIDE NAME]  |  Department of Computer Science & Engineering"
    pm2.font.size = Pt(13)
    pm2.font.color.rgb = RGBColor(200, 215, 235)
    pm2.space_before = Pt(6)
    
    pm3 = tf_m.add_paragraph()
    pm3.text = "Institution: [INSTITUTE NAME]  |  Academic Year 2026-2027"
    pm3.font.size = Pt(13)
    pm3.font.color.rgb = RGBColor(200, 215, 235)
    pm3.space_before = Pt(4)
    
    # -------------------------------------------------------------
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 2: Problem Statement...")
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, RGBColor(255, 255, 255))
    add_header(s2, "Problem Statement & Operational Motivation", "1. Introduction")
    
    # 3 Cards
    add_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8))
    tb2_1 = s2.shapes.add_textbox(Inches(0.95), Inches(1.95), Inches(3.3), Inches(4.5))
    tf2_1 = tb2_1.text_frame
    tf2_1.word_wrap = True
    p = tf2_1.paragraphs[0]
    p.text = "The Meeting Bottleneck"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    bullets1 = [
        "Knowledge workers spend 30-40% of working hours in multi-party meetings.",
        "Manual minutes taking is cognitively exhausting, subjective, and slow.",
        "Up to 50% of verbal action items and deadlines are forgotten or delayed without automated tracking.",
        "Meeting transcripts often exceed 20,000+ words, overwhelming participants."
    ]
    for b in bullets1:
        p = tf2_1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11.5)
        p.space_before = Pt(8)
        p.font.color.rgb = RGBColor(60, 60, 60)
        
    add_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8))
    tb2_2 = s2.shapes.add_textbox(Inches(4.95), Inches(1.95), Inches(3.3), Inches(4.5))
    tf2_2 = tb2_2.text_frame
    tf2_2.word_wrap = True
    p = tf2_2.paragraphs[0]
    p.text = "Core NLP Challenges"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    bullets2 = [
        "Acoustic Disfluencies: Spoken talk has false starts, fillers ('uh', 'um'), and repetitions.",
        "Severe Class Imbalance: Action items comprise only 5% to 35% of total conversational dialogue.",
        "Implicit Commitments: Natural agreements use varied phrasing rather than strict templates.",
        "Cascading Error: Speech recognition errors propagate directly into classifiers."
    ]
    for b in bullets2:
        p = tf2_2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11.5)
        p.space_before = Pt(8)
        p.font.color.rgb = RGBColor(60, 60, 60)
        
    add_card(s2, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8))
    tb2_3 = s2.shapes.add_textbox(Inches(8.95), Inches(1.95), Inches(3.4), Inches(4.5))
    tf2_3 = tb2_3.text_frame
    tf2_3.word_wrap = True
    p = tf2_3.paragraphs[0]
    p.text = "Proposed Solution"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    bullets3 = [
        "End-to-End Pipeline: Ingests raw audio, outputs transcripts, structured action tables, and minutes.",
        "Rigorous Benchmark: Evaluates 5 diverse ML & Transformer models on real municipal data.",
        "Honest Generalization: Employs an isolated 200-utterance human-verified ground-truth test set.",
        "High Efficiency: Executes natively on consumer GPU hardware with FP16 acceleration."
    ]
    for b in bullets3:
        p = tf2_3.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11.5)
        p.space_before = Pt(8)
        p.font.color.rgb = RGBColor(60, 60, 60)
        
    # -------------------------------------------------------------
    # SLIDE 3: SYSTEM ARCHITECTURE
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 3: Architecture...")
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, RGBColor(255, 255, 255))
    add_header(s3, "End-to-End Neural Pipeline Architecture", "2. System Design")
    
    stages = [
        ("Stage 1: ASR Engine", "OpenAI Whisper (base.en)", "• 16 kHz Mono Audio Ingestion\n• 80-channel log-mel filterbanks\n• Seq2seq Transformer Decoder\n• Native Windows FFmpeg Bypass", Inches(0.8)),
        ("Stage 2: Preprocessing", "Disfluency & Modals", "• Prunes fillers ('uh', 'um', 'like')\n• Retains modal auxiliaries\n• Contraction normalization\n• Lemmatization over Stemming", Inches(3.25)),
        ("Stage 3: Classifier", "5 Comparative Models", "• Logistic Regression (L2)\n• Linear Support Vector Machine\n• BiLSTM Recurrent Network\n• DistilBERT & RoBERTa", Inches(5.7)),
        ("Stage 4: Entity Extractor", "spaCy Dependency & NER", "• Head verb subtree traversal\n• Entity recognition (PERSON, ORG)\n• Temporal resolution (DATE, TIME)\n• Structured Task Triples", Inches(8.15)),
        ("Stage 5: Summarizer", "BART-large-CNN", "• Hierarchical Map-Reduce\n• 800-token chunk partitions\n• Multi-chunk intermediate maps\n• Executive Meeting Minutes", Inches(10.6))
    ]
    
    for title, sub, desc, left in stages:
        add_card(s3, left, Inches(2.0), Inches(2.3), Inches(4.6))
        tb = s3.shapes.add_textbox(left + Inches(0.1), Inches(2.1), Inches(2.1), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = RGBColor(24, 43, 73)
        
        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.size = Pt(10.5)
        p_s.font.bold = True
        p_s.font.color.rgb = RGBColor(31, 78, 121)
        p_s.space_before = Pt(3)
        
        for line in desc.split("\n"):
            p_d = tf.add_paragraph()
            p_d.text = line
            p_d.font.size = Pt(9.5)
            p_d.font.color.rgb = RGBColor(50, 50, 50)
            p_d.space_before = Pt(4)
            
    # -------------------------------------------------------------
    # SLIDE 4: DATASET & EXPERIMENTAL RIGOR
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 4: Dataset & Rigor...")
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, RGBColor(255, 255, 255))
    add_header(s4, "Dataset Selection & Strict Zero-Leakage Splitting", "3. Data Strategy")
    
    # Left Card
    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb4_1 = s4.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf4_1 = tb4_1.text_frame
    tf4_1.word_wrap = True
    p = tf4_1.paragraphs[0]
    p.text = "MeetingBank Corpus (Hugging Face)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    
    bullets4_1 = [
        "Corpus Source: huuuyeah/meetingbank derived from 1,366 city council and committee meetings across 6 major US municipal assemblies.",
        "Unconstrained Natural Discourse: High-stakes policy, municipal budget debates, public comments, and administrative motions.",
        "Strict Meeting-Level Partitioning: Zero random utterance shuffling to guarantee 0% data leakage across speakers, topics, and sessions:",
        "  - Training Set: 8 complete meetings (9,721 dialogue turns)",
        "  - Validation Set: 2 complete meetings (3,341 dialogue turns)",
        "  - Weak Test Set: 2 complete meetings (1,405 dialogue turns)"
    ]
    for b in bullets4_1:
        p = tf4_1.add_paragraph()
        p.text = "• " + b if not b.startswith("  -") else b
        p.font.size = Pt(11)
        p.space_before = Pt(6)
        p.font.color.rgb = RGBColor(50, 50, 50)
        
    # Right Card
    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    tb4_2 = s4.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf4_2 = tb4_2.text_frame
    tf4_2.word_wrap = True
    p = tf4_2.paragraphs[0]
    p.text = "Supervision & Ground-Truth Benchmark"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    
    bullets4_2 = [
        "Rule-Based Weak Labeler: Implemented via spaCy Matcher and regular expressions across 4 linguistic categories:",
        "  1. Commissive Intent ('I will', 'we shall')",
        "  2. Deontic Necessity ('we need to', 'must ensure')",
        "  3. Directives & Requests ('please make sure', 'can you')",
        "  4. Parliamentary Motions ('I move that', 'second the motion')",
        "The 200-Utterance Human-Verified Test Set:",
        "  - Independently audited and verified by human consensus.",
        "  - Balanced 51% / 49% distribution (98 action items, 102 negatives).",
        "  - Purpose: Honestly expose whether models memorize regex patterns or truly comprehend conversational semantics."
    ]
    for b in bullets4_2:
        p = tf4_2.add_paragraph()
        p.text = "• " + b if not b.startswith("  ") else b
        p.font.size = Pt(10.5)
        p.space_before = Pt(5)
        p.font.color.rgb = RGBColor(50, 50, 50)
        
    # -------------------------------------------------------------
    # SLIDE 5: PREPROCESSING & STEMMING VS LEMMATIZATION
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 5: Preprocessing...")
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, RGBColor(255, 255, 255))
    add_header(s5, "Data Preprocessing & Morphological Normalization", "4. Text Normalization")
    
    # Left: Explanation Card
    add_card(s5, Inches(0.8), Inches(1.8), Inches(5.4), Inches(4.8))
    tb5 = s5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.0), Inches(4.4))
    tf5 = tb5.text_frame
    tf5.word_wrap = True
    p = tf5.paragraphs[0]
    p.text = "Linguistic Preprocessing Nuances"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    
    bullets5 = [
        "Disfluency Pruning: Strips hesitations ('uh', 'um'), fillers ('you know', 'actually'), and stutter repetitions without altering core syntax.",
        "Custom Stopword Filtering: Standard stopword filters aggressively purge modal auxiliaries ('will', 'should', 'need', 'must'). In meeting dialogue, modals are the primary carriers of obligation. Our pipeline strictly preserves all modal auxiliary verbs.",
        "Morphological Normalization Trade-Off: Benchmarked NLTK Porter Stemmer against spaCy Lemmatizer across 1,000 utterances.",
        "Decision: spaCy Lemmatization preserves dictionary headwords essential for syntactic dependency parsing."
    ]
    for b in bullets5:
        p = tf5.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.space_before = Pt(8)
        p.font.color.rgb = RGBColor(50, 50, 50)
        
    # Right: Empirical Table Card
    add_card(s5, Inches(6.6), Inches(1.8), Inches(5.9), Inches(4.8))
    tb5_t = s5.shapes.add_textbox(Inches(6.8), Inches(2.0), Inches(5.5), Inches(0.5))
    p = tb5_t.text_frame.paragraphs[0]
    p.text = "Empirical Benchmark: Stemmer vs Lemmatizer"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    
    stem_df = pd.read_csv(os.path.join(res_dir, "preprocessing_stem_vs_lemma.csv"))
    add_table_to_slide(s5, stem_df, Inches(6.8), Inches(2.6), Inches(5.5), Inches(1.8), [1.3, 0.8, 0.9, 0.8, 0.8, 0.9])
    
    tb5_obs = s5.shapes.add_textbox(Inches(6.8), Inches(4.7), Inches(5.5), Inches(1.8))
    tf_obs = tb5_obs.text_frame
    tf_obs.word_wrap = True
    p = tf_obs.paragraphs[0]
    p.text = "Key Empirical Takeaway:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(31, 78, 121)
    p2 = tf_obs.add_paragraph()
    p2.text = "Porter Stemmer is ~11x faster (2,706 vs 247 samples/sec), but truncates valid roots ('scheduling' -> 'schedul'). spaCy Lemmatizer provides grammatically intact lemmas ('schedule') necessary for accurate dependency parsing."
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = RGBColor(60, 60, 60)
    p2.space_before = Pt(4)
    
    # -------------------------------------------------------------
    # SLIDE 6: EXPLORATORY DATA ANALYSIS
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 6: EDA...")
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, RGBColor(255, 255, 255))
    add_header(s6, "Exploratory Data Analysis: Key Findings", "5. Corpus Insights")
    
    # 4 Quadrants
    eda_items = [
        ("1. Severe Class Imbalance", "Action Items comprise 34.2% (3,322 turns) vs Non-Action 65.8% (6,399 turns). Requires class weighting and Macro-F1 optimization.", Inches(0.8), Inches(1.8)),
        ("2. Utterance Length Dynamics", "Action Items have median length of 22 words (mean: 28.4) vs Non-Action turns median of 8 words. Action commitments require syntactic complexity.", Inches(6.8), Inches(1.8)),
        ("3. Distinct Collocations", "Action items feature commitment bigrams ('make sure', 'need to', 'look forward'). Non-action is dominated by pleasantries ('thank you', 'roll call').", Inches(0.8), Inches(4.3)),
        ("4. Temporal U-Curve", "Action items concentrate in opening 15% (formal motions) and closing 20% (task directives) of meetings, while the middle is deliberative debate.", Inches(6.8), Inches(4.3))
    ]
    for title, desc, left, top in eda_items:
        add_card(s6, left, top, Inches(5.7), Inches(2.3))
        tb = s6.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.3), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = RGBColor(24, 43, 73)
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(60, 60, 60)
        p2.space_before = Pt(6)
        
    # -------------------------------------------------------------
    # SLIDE 7: FEATURE ENGINEERING & CHI-SQUARE
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 7: Feature Engineering...")
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, RGBColor(255, 255, 255))
    add_header(s7, "Feature Engineering & Statistical Feature Selection", "6. Features")
    
    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.4), Inches(4.8))
    tb7 = s7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.0), Inches(4.4))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    p = tf7.paragraphs[0]
    p.text = "Hybrid Feature Representation"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    bullets7 = [
        "Sparse TF-IDF Matrix (5,000 dimensions): Unigrams and bigrams with sublinear TF scaling (1 + log(tf)) and smoothed IDF.",
        "6 Dense Domain Linguistic Features:",
        "  1. Modal Auxiliary Verb Count ('will', 'need', 'must')",
        "  2. Imperative Sentence Head Indicator",
        "  3. Temporal Marker Count ('tomorrow', 'by friday')",
        "  4. Log-transformed Utterance Token Length",
        "  5. Sentiment Polarity / Decisiveness Score",
        "  6. First-Person Commissive Pronoun Count ('I', 'we')",
        "Concatenated Feature Matrix: Dimension (N, 5006)."
    ]
    for b in bullets7:
        p = tf7.add_paragraph()
        p.text = "• " + b if not b.startswith("  ") else b
        p.font.size = Pt(10.5)
        p.space_before = Pt(5)
        p.font.color.rgb = RGBColor(50, 50, 50)
        
    # Right: Chi-Square Table
    add_card(s7, Inches(6.6), Inches(1.8), Inches(5.9), Inches(4.8))
    tb7_t = s7.shapes.add_textbox(Inches(6.8), Inches(2.0), Inches(5.5), Inches(0.4))
    p = tb7_t.text_frame.paragraphs[0]
    p.text = "Top 10 Features by Chi-Square (χ²) Association"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    
    chi2_top10 = pd.read_csv(os.path.join(res_dir, "chi2_top_features.csv")).head(10)
    add_table_to_slide(s7, chi2_top10, Inches(6.8), Inches(2.5), Inches(5.5), Inches(3.8), [2.0, 1.8, 1.7])
    
    # -------------------------------------------------------------
    # SLIDE 8: TRADITIONAL ML & BILSTM
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 8: Traditional & BiLSTM...")
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, RGBColor(255, 255, 255))
    add_header(s8, "Model Development: Classical ML & BiLSTM", "7. Classifiers")
    
    cards_m = [
        ("Logistic Regression", "• Linear decision boundary with sigmoid\n• L2 regularization (C=1.0)\n• Balanced class weights\n• High recall (75.54%)\n• Validation Macro-F1: 74.78%", Inches(0.8)),
        ("Linear SVM", "• Soft-margin hyperplane (C=0.5)\n• L2-penalized squared hinge loss\n• High precision on weak data (72.31%)\n• Conservative decision boundary\n• Validation Macro-F1: 72.71%", Inches(4.8)),
        ("Bidirectional LSTM", "• Embedding(10000, 128)\n• BiLSTM(64) -> 128 hidden state\n• Dropout(0.3) -> Dense(32) -> Dense(1)\n• Weighted BCE Loss (pos_weight=1.93)\n• Validation Macro-F1: 75.84%", Inches(8.8))
    ]
    for title, desc, left in cards_m:
        add_card(s8, left, Inches(1.8), Inches(3.7), Inches(4.8))
        tb = s8.shapes.add_textbox(left + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(24, 43, 73)
        for line in desc.split("\n"):
            p_d = tf.add_paragraph()
            p_d.text = line
            p_d.font.size = Pt(11)
            p_d.font.color.rgb = RGBColor(50, 50, 50)
            p_d.space_before = Pt(6)
            
    # -------------------------------------------------------------
    # SLIDE 9: TRANSFORMER FINE-TUNING (DISTILBERT & ROBERTA)
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 9: Transformers...")
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, RGBColor(255, 255, 255))
    add_header(s9, "Contextual Transformer Fine-Tuning: DistilBERT & RoBERTa", "8. Transformers")
    
    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb9_1 = s9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf9_1 = tb9_1.text_frame
    tf9_1.word_wrap = True
    p = tf9_1.paragraphs[0]
    p.text = "DistilBERT (distilbert-base-uncased)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    bullets9_1 = [
        "Architecture: 6 transformer layers, 768 hidden dimension, 12 attention heads (66M parameters).",
        "Knowledge Distillation: Distilled from BERT-base; 40% smaller and 60% faster inference.",
        "Fine-Tuning: 3 epochs, AdamW optimizer, lr = 2e-5, linear warmup, max length 128.",
        "Weak Test Performance: 96.34% Macro-F1 (96.85% Precision, 93.29% Recall)."
    ]
    for b in bullets9_1:
        p = tf9_1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.space_before = Pt(8)
        p.font.color.rgb = RGBColor(50, 50, 50)
        
    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    tb9_2 = s9.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf9_2 = tb9_2.text_frame
    tf9_2.word_wrap = True
    p = tf9_2.paragraphs[0]
    p.text = "RoBERTa (roberta-base)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    bullets9_2 = [
        "Architecture: 12 transformer layers, 768 hidden dimension, 12 attention heads (125M parameters).",
        "Key Refinements: Dynamic token masking, removed NSP loss, byte-level BPE, 160 GB pre-training corpus.",
        "Hardware Optimization: Fine-tuned on NVIDIA RTX 3050 GPU using FP16 mixed precision and gradient accumulation (eff. batch 16).",
        "Benchmark Champion: 94.48% Macro-F1 on human-verified ground-truth test set."
    ]
    for b in bullets9_2:
        p = tf9_2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.space_before = Pt(8)
        p.font.color.rgb = RGBColor(50, 50, 50)
        
    # -------------------------------------------------------------
    # SLIDE 10: MASTER RESULTS LEADERBOARD
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 10: Master Leaderboard...")
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, RGBColor(255, 255, 255))
    add_header(s10, "Empirical Results: Master 5-Model Comparison", "9. Evaluation")
    
    # Leaderboard table
    master_df = pd.read_csv(os.path.join(res_dir, "all_5_models_master_comparison.csv"))
    # Filter to human verified test set for primary highlight
    hv_df = master_df[master_df['Split'] == 'Human-Verified Test'][['Model', 'Accuracy', 'Precision', 'Recall', 'Action_Item_F1', 'Macro_F1', 'ROC_AUC']].copy()
    
    add_card(s10, Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.6))
    tb10_t = s10.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(0.35))
    p = tb10_t.text_frame.paragraphs[0]
    p.text = "Ground-Truth Leaderboard (200 Human-Verified Utterances Benchmark)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    
    add_table_to_slide(s10, hv_df, Inches(1.0), Inches(2.35), Inches(11.3), Inches(1.9), [2.3, 1.5, 1.5, 1.5, 1.5, 1.5, 1.5])
    
    # Bottom Comparison Cards
    add_card(s10, Inches(0.8), Inches(4.6), Inches(5.7), Inches(2.2))
    tb10_c1 = s10.shapes.add_textbox(Inches(1.0), Inches(4.7), Inches(5.3), Inches(2.0))
    tf10_c1 = tb10_c1.text_frame
    tf10_c1.word_wrap = True
    p = tf10_c1.paragraphs[0]
    p.text = "Traditional ML Generalization Collapse"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(160, 40, 40)
    p2 = tf10_c1.add_paragraph()
    p2.text = "Linear SVM suffered an acute recall drop to 54.08% on the human-verified benchmark (Macro-F1: 71.41%). Classical models memorized the specific lexical triggers of the weak supervision regexes and failed on natural conversational paraphrases."
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = RGBColor(60, 60, 60)
    p2.space_before = Pt(4)
    
    add_card(s10, Inches(6.8), Inches(4.6), Inches(5.7), Inches(2.2))
    tb10_c2 = s10.shapes.add_textbox(Inches(7.0), Inches(4.7), Inches(5.3), Inches(2.0))
    tf10_c2 = tb10_c2.text_frame
    tf10_c2.word_wrap = True
    p = tf10_c2.paragraphs[0]
    p.text = "Transformer Semantic Superiority"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(20, 100, 40)
    p2 = tf10_c2.add_paragraph()
    p2.text = "RoBERTa established state-of-the-art performance with 94.48% Macro-F1 (97.80% Precision, 90.82% Recall). DistilBERT achieved 93.48% Macro-F1 with 60% faster inference. Pre-trained attention captures genuine commitment semantics."
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = RGBColor(60, 60, 60)
    p2.space_before = Pt(4)
    
    # -------------------------------------------------------------
    # SLIDE 11: ASR, NER, SUMMARIZATION & DEMO
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 11: End-to-End Pipeline...")
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, RGBColor(255, 255, 255))
    add_header(s11, "ASR, Information Extraction, Summarization & Demo", "10. Full System")
    
    cards_sys = [
        ("Whisper ASR (base.en)", "• 8 Speech audio clips evaluated\n• Average WER: 3.78%\n• Average CER: 0.96%\n• Perfect 0.00% WER on 5 clips\n• Robust to acoustic noise", Inches(0.8)),
        ("spaCy IE & NER", "• 98 Action Items extracted\n• Task: Verb subtree dependency\n• Owner: Named entity & pronouns\n• Deadline: Temporal dates / times\n• Structured tabular output", Inches(4.8)),
        ("BART Summarization", "• facebook/bart-large-cnn\n• Hierarchical Map-Reduce chunking\n• Mean ROUGE-1 F1: 0.1334\n• Mean ROUGE-L F1: 0.1072\n• Cohesive executive minutes", Inches(8.8))
    ]
    for title, desc, left in cards_sys:
        add_card(s11, left, Inches(1.8), Inches(3.7), Inches(2.6))
        tb = s11.shapes.add_textbox(left + Inches(0.2), Inches(1.95), Inches(3.3), Inches(2.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = RGBColor(24, 43, 73)
        for line in desc.split("\n"):
            p_d = tf.add_paragraph()
            p_d.text = line
            p_d.font.size = Pt(10)
            p_d.font.color.rgb = RGBColor(50, 50, 50)
            p_d.space_before = Pt(3)
            
    # Bottom Demo Box
    add_card(s11, Inches(0.8), Inches(4.6), Inches(11.7), Inches(2.2), bg_rgb=RGBColor(240, 245, 252))
    tb_demo = s11.shapes.add_textbox(Inches(1.0), Inches(4.7), Inches(11.3), Inches(2.0))
    tf_demo = tb_demo.text_frame
    tf_demo.word_wrap = True
    p = tf_demo.paragraphs[0]
    p.text = "Live End-to-End Execution Sample (demo.py)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(31, 78, 121)
    
    demo_lines = [
        "• Audio Input: 'We need to finalize the quarterly budget review by Friday afternoon.'",
        "• RoBERTa Classifier: Action Item Identified (Confidence: 99.90%)",
        "• Information Extraction: Task = 'finalize the quarterly budget review' | Owner = 'Joint Team' | Deadline = 'Friday, afternoon'",
        "• Executive Summary: 'The city finance director noted the quarterly budget review is scheduled for completion and release Friday.'"
    ]
    for dl in demo_lines:
        p = tf_demo.add_paragraph()
        p.text = dl
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(40, 40, 40)
        p.space_before = Pt(2)
        
    # -------------------------------------------------------------
    # SLIDE 12: CONCLUSION & FUTURE SCOPE
    # -------------------------------------------------------------
    print("[Presentation Generator] Generating Slide 12: Conclusion & Future Scope...")
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, RGBColor(24, 43, 73)) # Deep Navy
    
    add_card(s12, Inches(0.8), Inches(1.0), Inches(11.733), Inches(5.5),
             bg_rgb=RGBColor(31, 56, 94), border_rgb=RGBColor(60, 95, 145))
             
    tb12 = s12.shapes.add_textbox(Inches(1.2), Inches(1.3), Inches(10.9), Inches(4.8))
    tf12 = tb12.text_frame
    tf12.word_wrap = True
    
    p = tf12.paragraphs[0]
    p.text = "Conclusion & Future Research Horizons"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    bullets12 = [
        "Completed Rigorous Pipeline: Successfully engineered, evaluated, and documented an end-to-end meeting transcription and action-item summarizer.",
        "Empirical Benchmark Champion: Fine-tuned RoBERTa achieved 94.48% Macro-F1 on human-verified ground-truth, proving transformer semantic robustness over classical linear models.",
        "Reproducibility: Fully trained on NVIDIA RTX 3050 GPU; complete code, models, figures, and datasets preserved in repository.",
        "Future Work 1 (Streaming ASR): Transition to low-latency real-time streaming speech transcription for live video meetings.",
        "Future Work 2 (Speaker Diarization): Integrate pyannote-audio to identify speaker turn boundaries from raw multi-speaker room audio.",
        "Future Work 3 (LLMs & RAG): Unify classification and extraction with quantized open-weights LLMs (LLaMA-3 / Mistral) paired with vector database memory."
    ]
    for b in bullets12:
        p = tf12.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(230, 240, 255)
        p.space_before = Pt(8)
        
    print(f"[Presentation Generator] Saving presentation to {out_path}...")
    prs.save(out_path)
    print(f"[Presentation Generator] Presentation successfully created at {out_path}!")

if __name__ == "__main__":
    build_presentation()
